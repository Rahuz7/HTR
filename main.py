import sys, os
import cv2


from PyQt5.QtWidgets import QMainWindow, QApplication, QWidget, QPushButton, QFileDialog, QLabel, QLineEdit, QProgressBar, QPlainTextEdit
from PyQt5.QtGui import QIcon, QPixmap
from PyQt5.QtCore import QBasicTimer
from docx import Document
from fpdf import FPDF
from docx.shared import Inches
import shutil

from segmentation import TextSegmentation
from preprocessing import Preprocessing
from classification import classify



class Window(QMainWindow):

    def __init__(self):
        QMainWindow.__init__(self)
        self.title = 'TextRecognition'
        self.left = 10
        self.top = 10
        self.width = 1112
        self.height = 923

        self.stylesheet = """

        QMainWindow{
            background-color: #f4f4f4;
        }

        QPushButton{
            background-color: #f2f2f2;
            border: 1px solid green;
            border-radius: 5px;
        }

        QLineEdit{
            border: 1px solid;
        }

        QPushButton#pdfButton{
            background-image: url("./assets/icons/pdf.png");
        }

        QPushButton#wordButton{
            background-image: url("./assets/icons/word.png");
        }

        QPushButton#txtButton{
            background-image: url("./assets/icons/txt.png");
        }

        QLabel#title{
            font-size: 20px;
            padding: 10px;
            text-align: center;
            font-family: Helvetica;
            color: green;
        }

        QProgressBar{
            height: 50px;
            width: 200px;
            margin: 20px;
        }
        """
        self.initUI()

    def initUI(self):
        self.setWindowTitle(self.title)
        self.setGeometry(self.left, self.top, self.width, self.height)
        self.setStyleSheet(self.stylesheet)

        self.name = QLabel(self)
        self.name.setText("Handwritten Text Recognition")
        self.name.move(450, 84)
        self.name.resize(406, 70)
        self.name.setObjectName("title")
        self.name.setStyleSheet(self.stylesheet)

        self.inputBox = QLineEdit(self)
        self.inputBox.move(64, 320)
        self.inputBox.resize(442, 39)
        self.inputBox.setReadOnly(True)

        self.inputButton = QPushButton("Select a file", self)
        self.inputButton.resize(150, 39)
        self.inputButton.move(195, 370)
        self.inputButton.setToolTip("Click here to choose an image")
        self.inputButton.setStyleSheet(self.stylesheet)
        self.inputButton.clicked.connect(self.openFile)

        self.outputBox = QLineEdit(self)
        self.outputBox.move(64, 490)
        self.outputBox.resize(442, 39)
        self.outputBox.setReadOnly(True)

        self.outputButton = QPushButton("Save destination", self)
        self.outputButton.resize(150, 39)
        self.outputButton.move(195, 540)
        self.outputButton.setToolTip("Click here to choose the save destination of the output")
        self.outputButton.clicked.connect(self.saveFile)
        self.outputButton.setStyleSheet(self.stylesheet)

        self.startButton = QPushButton("Start", self)
        self.startButton.resize(191, 39)
        self.startButton.move(450, 780)
        self.startButton.clicked.connect(self.doAction)
        self.startButton.setToolTip("Click to start")
        self.startButton.setStyleSheet(self.stylesheet)

        self.pdfButton = QPushButton("", self)
        self.pdfButton.resize(180, 180)
        self.pdfButton.move(815, 189)
        self.pdfButton.setToolTip("Click here to output as PDF")
        self.pdfButton.clicked.connect(self.save_to_pdf)
        self.pdfButton.setStyleSheet(self.stylesheet)
        self.pdfButton.setObjectName("pdfButton")

        self.wordButton = QPushButton("", self)
        self.wordButton.resize(180, 180)
        self.wordButton.move(815, 390)
        self.wordButton.setToolTip("Click here to output as Word")
        self.wordButton.clicked.connect(self.save_to_word)
        self.wordButton.setStyleSheet(self.stylesheet)
        self.wordButton.setObjectName("wordButton")

        self.txtButton = QPushButton("", self)
        self.txtButton.resize(180, 180)
        self.txtButton.move(815, 591)
        self.txtButton.setToolTip("Click here to output as txt")
        self.txtButton.clicked.connect(self.save_to_txt)
        self.txtButton.setStyleSheet(self.stylesheet)
        self.txtButton.setObjectName("txtButton")

        self.progressBar = QProgressBar(self)
        self.progressBar.setGeometry(262, 808, 589, 80)
        self.progressBar.setObjectName("progressBar")

        self.timer = QBasicTimer()
        self.step = 0

        self.show()

    def openFile(self):
        options = QFileDialog.Options()
        options |= QFileDialog.DontUseNativeDialog
        self.img, _ = QFileDialog.getOpenFileName(self,"QFileDialog.getOpenFileName()", "","All Files (*)", options=options) #;;Python Files (*.png)
        self.inputBox.setText(self.img)
        self.img = cv2.imread(self.img)

    def saveFile(self):
        options = QFileDialog.Options()
        options |= QFileDialog.DontUseNativeDialog
        self.save_destination = str(QFileDialog.getExistingDirectory(self, "Select Directory"))
        self.outputBox.setText(self.save_destination)

    def main(self):
        print("started")
        p = Preprocessing()
        resized = p.resize(self.img)
        denoised = p.denoise(resized)
        s = TextSegmentation(denoised)
        s.linesSegmentation()
        self.result = classify()
        shutil.rmtree("./out")
        print("Finished")


    def start(self):
        pass

    def timerEvent(self, e):
        if self.step >= 100:
            self.timer.stop()
            self.startButton.setText('Finished')
            return

        self.step = self.step + 1
        self.progressBar.setValue(self.step)


    def doAction(self):
        if self.timer.isActive():
            self.timer.stop()
            self.startButton.setText('Start')
        else:
            self.timer.start(100, self)
            self.startButton.setText('Stop')

    def save_to_pdf(self):
        self.main()
        pdf = FPDF()
        pdf.add_page()
        pdf.set_xy(10, 10)
        pdf.set_font("arial", "B", 13.0)
        pdf.cell(ln = 0, h=5.0, align="L", w=0, txt=self.result, border=0)
        try:
            save_destination = self.save_destination + "/result.pdf"
            pdf.output(save_destination, "F")
        except:
            pdf.output("result.pdf", "F")

    def save_to_word(self):
        self.main()
        document = Document()
        document.add_paragraph(self.result)
        try:
            save_destination = self.save_destination + "/result.docx"
            document.save(save_destination)
        except:
            document.save("result.docx")

    def save_to_txt(self):
        self.main()
        try:
            save_destination = self.save_destination + "/result.txt"
            f = open(save_destination, "w+")
        except:
            f = open("result.txt", "w+")
        f.write(self.result)
        f.close()

if __name__ == "__main__":
    app = QApplication(sys.argv)
    fen = Window()
    app.exec_()
